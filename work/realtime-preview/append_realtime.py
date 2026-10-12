from pathlib import Path
from copy import deepcopy
from hashlib import sha256
from zipfile import ZipFile, ZIP_DEFLATED
from lxml import etree
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root = Path('/Users/blackwave/Documents/Crest')
work = root / 'work/realtime-preview'
src = root / 'output/Crest经营看板与订单页七版方案.docx'
out = root / 'output/Crest经营看板与订单页方案_含今日实时模拟.docx'
source_hash = sha256(src.read_bytes()).hexdigest()
doc = Document(src)
original_count = len(doc.element.body) - 1

def text(value, size=11, bold=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = Pt(16)
    r = p.add_run(value)
    r.bold = bold
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor(0, 0, 0)
    r.font.name = 'Arial'
    r._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Arial Unicode MS')
    return p

def heading(value, level=1):
    p = doc.add_heading(value, level)
    p.paragraph_format.space_before = Pt(10 if level == 2 else 0)
    p.paragraph_format.space_after = Pt(6)
    for r in p.runs:
        r.font.color.rgb = RGBColor(0, 0, 0)
        r._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Arial Unicode MS')
    return p

def picture(name, width):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1
    p.paragraph_format.space_after = Pt(0)
    p.add_run().add_picture(str(work / name), width=Inches(width))
    return p

doc.add_page_break()
heading('看板 E 今日实时')
text('完整页面 · 先读左列上半页，再读右列下半页 · 全部为模拟数据', 10)
picture('live-spread.png', 7.0)
text('页面自动更新，无暂停、倒计时或手动刷新区。横向门店图的完整内容见下一页。', 9)

doc.add_page_break()
heading('看板 E 设计说明')
heading('使用者与核心问题', 2)
text('适合营业中的老板、区域负责人和店长。先判断截至此刻的收入与昨日同时段相比如何，再找到需要关注的门店与渠道。')
heading('页面结构与图表交互', 2)
text('日期与门店范围 → 净实收、有效订单数、平均每单实收 → 累计净实收曲线 → 最多两家值得关注的门店 → 门店渠道堆叠柱图 → 渠道金额与占比。所有内容随门店选择一起变化。')
text('累计曲线以元为单位，对比 10/09 与 10/08 从营业日起点至同一截止时间的累计净实收。曲线便于判断营业进度；未到时段留空。轻点或拖动查看时间点，也可使用左右按钮。')
text('门店采用从零起算的纵向细柱，堂食紫色、外卖珊瑚红、自提青绿色，保留金额与较昨日同时段变化率。手机左右滑动查看 10 店；点柱或关注摘要，只展开该店的收入、订单、每单实收与渠道汇总。')
text('渠道以紧凑占比条配合金额和占比数值，快速识别收入结构。占比仅对当前完整覆盖的有效范围计算；不把占比高低直接解释为增长或下降原因。')
heading('具体判断情景', 2)
text('模拟 14:32，全店净实收 121,860 元，较昨日同时段增加 4,620 元（+3.9%），订单数 +1.7%、平均每单实收 +2.2%。但江湾店减少 3,000 元（−16.7%，上期 18,000 元）。负责人可优先查看江湾店汇总，核对订单与渠道差异后再跟进。')
heading('全部门店展开图', 2)
picture('live-wide.png', 6.8)
text('仅为文档展开横向区域；手机中保留左右滑动。所有柱体使用同一尺度。', 9)

doc.add_page_break()
heading('自动更新与数据接入')
heading('自动更新方式', 2)
text('模拟每 10 秒更新一次，核心数字、曲线、门店和渠道使用同一份数据同步变化。页面无需操作，不展示更新状态模块；截止时间保留在净实收卡片中。演示状态选择器位于手机外，仅用于预览不同状态。')
text('真实接入后的刷新频率需根据数据源延迟与接口能力确定。取到相同或较旧的数据时不伪造增量；页面展示数据实际截止时间。失败或延迟时保留最近有效结果，并在原有指标说明中提示。')
heading('状态设计', 2)
table = doc.add_table(rows=1, cols=3)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
widths = [0.9, 2.8, 3.3]
for col,w in zip(table.columns,widths): col.width=Inches(w)
for c,w,t in zip(table.rows[0].cells,widths,['状态','页面表现','统计与更新规则']):
    c.width = Inches(w); c.text = t
rows = [
    ('正常','数据自动变化，不显示刷新控制。','各区域使用同一截止时间和门店范围。'),
    ('数据延迟','说明中提示延迟，显示原截止时间。','保留最近有效结果，不以请求成功时间冒充数据时间。'),
    ('部分缺失','提示缺数门店，柱图标注缺数。','从本期、上期同时排除该门店；不把缺数当零。'),
    ('请求失败','说明中提示暂未取得新数据。','保留最近结果；真实接入时后台重试。'),
    ('首次加载','指标显示占位，图表显示加载提示。','没有有效快照前不显示数值或比较结果。'),
    ('暂无数据','指标显示横线，图表说明暂无数据。','不生成零收入、下降率或渠道占比。'),
]
for row in rows:
    cells = table.add_row().cells
    for c,w,t in zip(cells,widths,row): c.width=Inches(w); c.text=t
for ri,row in enumerate(table.rows):
    for ci,c in enumerate(row.cells):
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        tcPr=c._tc.get_or_add_tcPr()
        margins=OxmlElement('w:tcMar')
        for edge,val in [('top',90),('left',100),('bottom',90),('right',100)]:
            e=OxmlElement('w:'+edge);e.set(qn('w:w'),str(val));e.set(qn('w:type'),'dxa');margins.append(e)
        tcPr.append(margins)
        borders=OxmlElement('w:tcBorders')
        for edge in ['top','left','bottom','right']:
            e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
        tcPr.append(borders)
        if ri==0:
            fill=OxmlElement('w:shd');fill.set(qn('w:fill'),'E9EEF5');tcPr.append(fill)
        for p in c.paragraphs:
            p.paragraph_format.line_spacing = Pt(14)
            p.paragraph_format.space_after = Pt(0)
            if ci==0: p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.font.size=Pt(10);r.bold=ri==0;r.font.color.rgb=RGBColor(0,0,0)
                r._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Arial Unicode MS')
heading('现有字段与接入前核对', 2)
text('已有净实收、订单数、门店与渠道汇总可支撑基础指标和结构图，具体字段仍需核对。平均每单实收按同口径总净实收 ÷ 总订单数计算，不平均门店均值。')
text('当日实时曲线与昨日同时段比较，还需要小时或分钟累计数据、营业日与时区定义、数据实际截止时间、各店完整性标记，以及可保证同一批次更新的数据版本。仅有按日趋势不能直接支持这些能力。')
text('上线前需确认有效订单、退款与取消订单口径，渠道是否互斥且覆盖完整总额，以及新开、停业和缺数门店的可比范围。上期为零或缺失时不计算增长百分比。当前模拟未连接真实接口。')

tmp = work / 'authored.docx'
doc.save(tmp)

# Preserve every original package part byte for byte except the document body
# and the additions required for the two new images.
with ZipFile(src) as before, ZipFile(tmp) as authored:
    original_xml = etree.fromstring(before.read('word/document.xml'))
    new_xml = etree.fromstring(authored.read('word/document.xml'))
    ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    original_body=original_xml.find('w:body',ns)
    new_body=new_xml.find('w:body',ns)
    for a,b in zip(list(original_body)[:-1],list(new_body)[:original_count]):
        assert etree.tostring(a,method='c14n')==etree.tostring(b,method='c14n'), 'Original content changed'
    additions=list(new_body)[original_count:-1]
    for node in additions: original_body.insert(len(original_body)-1,deepcopy(node))
    changes={'word/document.xml':etree.tostring(original_xml,xml_declaration=True,encoding='UTF-8',standalone=True)}
    # Retain the original relationships, appending only new image relations.
    relpath='word/_rels/document.xml.rels'
    relold=etree.fromstring(before.read(relpath));relnew=etree.fromstring(authored.read(relpath))
    ids={n.get('Id') for n in relold}
    for r in relnew:
        if r.get('Id') not in ids: relold.append(deepcopy(r))
    changes[relpath]=etree.tostring(relold,xml_declaration=True,encoding='UTF-8',standalone=True)
    with ZipFile(out,'w',ZIP_DEFLATED) as result:
        for info in before.infolist(): result.writestr(info,changes.get(info.filename,before.read(info.filename)))
        oldnames=set(before.namelist())
        for name in authored.namelist():
            if name not in oldnames: result.writestr(name,authored.read(name))
assert sha256(src.read_bytes()).hexdigest()==source_hash, 'Source changed during append'
with ZipFile(src) as before,ZipFile(out) as after:
    unchanged=[n for n in before.namelist() if n not in changes]
    assert all(before.read(n)==after.read(n) for n in unchanged)
check=Document(out)
assert len(check.inline_shapes)==len(Document(src).inline_shapes)+2
print(f'Saved {out.name}; preserved {original_count} original body nodes and {len(unchanged)} package parts; appended {len(additions)} nodes; {len(check.inline_shapes)} images.')
