from pathlib import Path
from copy import deepcopy
from PIL import Image
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

BASE=Path('/Users/blackwave/Documents/Crest/work/word-export')
OUT=Path('/Users/blackwave/Documents/Crest/output/Crest经营看板与订单页七版方案_完整高清版.docx')
doc=Document()
s=doc.sections[0]
s.page_width=Inches(8.5);s.page_height=Inches(11)
s.top_margin=s.bottom_margin=Inches(.65)
s.left_margin=s.right_margin=Inches(.65)
s.footer_distance=Inches(.3)

for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
 st=doc.styles[name];st.font.name='Arial';st.font.color.rgb=RGBColor(0,0,0)
 st._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Arial Unicode MS')
 st.paragraph_format.space_after=Pt(7)
doc.styles['Normal'].font.size=Pt(11)
doc.styles['Normal'].paragraph_format.line_spacing=Pt(16)
doc.styles['Normal'].paragraph_format.space_after=Pt(6)
doc.styles['Title'].font.size=Pt(26)
doc.styles['Heading 1'].font.size=Pt(21)
doc.styles['Heading 2'].font.size=Pt(12)
doc.styles['Heading 2'].paragraph_format.space_before=Pt(7)
doc.styles['Heading 2'].paragraph_format.space_after=Pt(3)
doc.styles['Heading 2'].paragraph_format.line_spacing=Pt(16)
doc.styles['Subtitle'].font.size=Pt(11)
doc.styles['Caption'].font.size=Pt(9)
doc.styles['Caption'].font.color.rgb=RGBColor.from_string('666666')
for st in doc.styles:
 for el in st._element.iter(qn('w:rFonts')):
  for attr in list(el.attrib):
   if 'theme' in attr.lower():del el.attrib[attr]
 for el in list(st._element.iter(qn('w:pBdr'))):el.getparent().remove(el)
 if st.type==1:
  sp=OxmlElement('w:snapToGrid');sp.set(qn('w:val'),'0');st._element.get_or_add_pPr().append(sp)
for el in list(s._sectPr.findall(qn('w:docGrid'))):s._sectPr.remove(el)

f=s.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=f.add_run('Crest 方案评审  ·  ');r.font.size=Pt(9);r.font.color.rgb=RGBColor.from_string('777777')
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
doc.core_properties.title='Crest 经营看板与订单页方案汇编'
doc.core_properties.subject='四版经营看板与三版订单汇总页面的手机模拟和设计说明'
doc.core_properties.author='Crest'

def p(text='',style=None):return doc.add_paragraph(text,style)
def h(text,level=2):return doc.add_heading(text,level)
def page():doc.add_page_break()

def table(headers, rows, widths):
 t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 for c,w in zip(t.columns,widths):c.width=Inches(w)
 pr=t._tbl.tblPr;b=OxmlElement('w:tblBorders')
 for k in ['top','left','bottom','right','insideH','insideV']:
  e=OxmlElement('w:'+k);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'5');e.set(qn('w:color'),'D9D9D9');b.append(e)
 pr.append(b)
 for idx,vals in enumerate([headers]+rows):
  cells=t.rows[0].cells if idx==0 else t.add_row().cells
  for ci,(c,text,w) in enumerate(zip(cells,vals,widths)):
   c.width=Inches(w);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   cp=c._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'30363D' if idx==0 else ('F3F5F7' if idx%2==0 else 'FFFFFF'));cp.append(sh)
   m=OxmlElement('w:tcMar')
   for k,v in [('top',95),('bottom',95),('left',110),('right',110)]:
    e=OxmlElement('w:'+k);e.set(qn('w:w'),str(v));e.set(qn('w:type'),'dxa');m.append(e)
   cp.append(m)
   pp=c.paragraphs[0];pp.paragraph_format.space_after=Pt(0);pp.paragraph_format.line_spacing=Pt(14)
   rr=pp.add_run(text);rr.font.size=Pt(10)
   if idx==0:rr.bold=True;rr.font.color.rgb=RGBColor(255,255,255)
   if ci==0:pp.alignment=WD_ALIGN_PARAGRAPH.CENTER
  tr=t.rows[idx]._tr.get_or_add_trPr();tr.append(OxmlElement('w:cantSplit'))
  if idx==0:tr.append(OxmlElement('w:tblHeader'))
 p().paragraph_format.space_after=Pt(0)
 return t

def floating_phone(name,alt):
 pp=p();pp.paragraph_format.space_after=Pt(0);pp.paragraph_format.line_spacing=Pt(1);pp.paragraph_format.space_before=Pt(0)
 rr=pp.add_run();shape=rr.add_picture(str(BASE/'images'/f'{name}.png'),width=Inches(3.18))
 inline=shape._inline
 anchor=OxmlElement('wp:anchor')
 for k,v in {'distT':'0','distB':'72000','distL':'0','distR':'190500','simplePos':'0','relativeHeight':'0','behindDoc':'0','locked':'0','layoutInCell':'1','allowOverlap':'0'}.items():anchor.set(k,v)
 pos=OxmlElement('wp:simplePos');pos.set('x','0');pos.set('y','0');anchor.append(pos)
 for axis,relative in [('H','column'),('V','paragraph')]:
  el=OxmlElement('wp:position'+axis);el.set('relativeFrom',relative);off=OxmlElement('wp:posOffset');off.text='0';el.append(off);anchor.append(el)
 anchor.append(deepcopy(inline.find(qn('wp:extent'))))
 ef=OxmlElement('wp:effectExtent')
 for k in ['l','t','r','b']:ef.set(k,'0')
 anchor.append(ef)
 wrap=OxmlElement('wp:wrapSquare');wrap.set('wrapText','right');anchor.append(wrap)
 dp=deepcopy(inline.find(qn('wp:docPr')));dp.set('descr',alt);anchor.append(dp)
 anchor.append(deepcopy(inline.find(qn('wp:cNvGraphicFramePr'))));anchor.append(deepcopy(inline.find(qn('a:graphic'))))
 inline.getparent().replace(inline,anchor)

def picture(path, width=6.8, maxheight=None, alt=''):
 im=Image.open(path)
 if maxheight: width=min(width,maxheight*im.width/im.height)
 pp=p();pp.alignment=WD_ALIGN_PARAGRAPH.CENTER
 pp.paragraph_format.line_spacing=1
 pp.paragraph_format.space_before=Pt(0);pp.paragraph_format.space_after=Pt(0)
 shape=pp.add_run().add_picture(str(path),width=Inches(width))
 shape._inline.docPr.set('descr',alt)

def variant(code,title,subtitle,image,sections):
 page();h(code+' '+title,1)
 p('完整页面 · 先读左列上半页，再读右列下半页 · 全部为模拟数据','Caption')
 picture(BASE/'full-images'/f'{image}-spread.png',6.8,8.35,title+'完整手机页面，左列为上半页、右列为下半页')
 page();h(code+' 设计说明',1);p(subtitle,'Subtitle')
 for label,text in sections:
  h(label);p(text)
 if image in ['dashboard-a','dashboard-b','order-b']:
  h('10 家门店完整展开')
  picture(BASE/'full-images'/f'{image}-wide.png',6.8,3.3,'横向滑动区域的完整门店图')
  p('为文档阅读展开横向区域；手机中仍按原设计左右滑动。','Caption')

h('Crest 经营看板与订单页方案汇编',0)
p('7 版手机端演示模拟 · 完整高清版  2026 年 10 月 9 日','Subtitle')
p('本文件整理 4 版经营看板与 3 版订单页，供页面方向评审和后续设计落地使用。建议以经营总览作为默认看板，以订单总览作为默认订单页；其余版本分别服务门店跟进、变化分析和周期观察。')
p('共同目标是让使用者在 10 秒内了解经营表现，在 30 秒内发现值得关注的门店或渠道，并知道下一步查看哪一层汇总。订单页面向每天数百至上千单的经营规模，只展示全局与分组统计。')
h('版本目录')
table(['编号','版本','主要问题','定位'],[
 ['看板 A','经营总览','整体表现如何，下一步看哪里','默认看板'],
 ['看板 B','门店关注','先跟进哪家门店，依据是什么','区域负责人'],
 ['看板 C','变化来源','波动伴随单量还是每单实收变化','经营分析'],
 ['看板 D','经营日历热力图','哪些日期集中繁忙或偏淡','周期观察'],
 ['订单 A','订单总览','整体单量及渠道结构如何变化','默认订单页'],
 ['订单 B','门店对比','哪些门店贡献了单量变化','门店汇总比较'],
 ['订单 C','渠道与退款','渠道单量与退款情况如何变化','交易质量观察'],
],[.8,1.55,3.15,1.7])
h('已确认的设计方向')
p('趋势图使用曲线折线图。门店图保留纵向细柱，支持横向滑动；净实收柱按堂食、外卖分段，存在自提时保留自提段。渠道颜色固定为堂食紫色、外卖珊瑚红、自提青绿色。')
p('订单模块不提供逐笔订单列表，不展示订单号、顾客或商品详情，也不引导定位到某一笔订单。点击行为均停留在日期、门店或渠道的汇总层级。')
h('阅读说明')
p('截图已按高分辨率重新导出，保留现有浅色或深色外观。每版完整页面沿卡片间隙分成左右两列，先读左列再读右列；门店横向图另附完整展开图。全部数值为模拟数据；各组模拟的数据集、周期与分母口径可能不同，不应用于跨版本经营数值比较。文中的字段支持判断仅依据目前已知的数据边界，仍需在落地前核对。')

variant('看板 A','经营总览','推荐默认看板  先判断整体表现，再找到关注方向','dashboard-a',[
 ('适用角色与问题','老板和店长。快速判断净实收是增长还是下降，以及应继续查看趋势、门店还是渠道。'),
 ('页面顺序','日期、门店与更新时间 → 净实收、订单数、平均每单实收 → 每日趋势 → 门店净实收 → 渠道构成 → 最多 3 条重点变化。'),
 ('图表与选择理由','每日净实收用曲线展示，单位元/日，可切换 7 天和 30 天，上期用虚线。门店用渠道堆叠细柱比较规模与结构；渠道占比条配金额，控制在 3 个主要图表内。'),
 ('点击后的行为','点选日期查看本期与对应上期数值；长按拖动连续查看。门店柱横向滑动，点击打开门店汇总；渠道入口查看该渠道金额、占比和门店分布。'),
 ('具体使用情景','示例净实收 16.8 万元，较上期少 1.2 万元；订单数持平。用户据此关注每单实收变化，再查看江湾店及堂食渠道，而非直接认定销量下降。'),
 ('数据支持与补充','已有净实收、订单数、按日及门店渠道汇总可支撑主体。需确认有效订单分母、退款归属、营业日边界和完整性标记。')
])
variant('看板 B','门店关注','适合区域负责人  以减少金额确定跟进顺序','dashboard-b',[
 ('适用角色与问题','区域负责人和多店老板。先判断哪家门店影响最大，再核对下降幅度及上期基数。'),
 ('页面顺序','日期与门店范围 → 整体净实收及差异 → 门店对比图 → 选中门店摘要 → 重点关注门店。'),
 ('图表与选择理由','主图为纵向细柱，可切换增减金额、变化率和净实收。默认按减少金额排序；增减图保留零基线，净实收从零起算并按渠道堆叠。单位随模式显示元或百分比。'),
 ('点击后的行为','横滑浏览 10 家门店；点柱更新门店摘要，进入该店订单、每单实收及渠道汇总。可切换为只查看这家店，再用相同筛选查看日趋势。'),
 ('具体使用情景','江湾店减少 6,000 元、下降 20%；文汇店下降 50%，但仅减少 1,200 元。负责人先跟进江湾，兼顾文汇的低基数，避免只被最大百分比吸引。'),
 ('数据支持与补充','门店本期和上期净实收可直接支持对比。需补充新开、停业和缺数标记，确认可比门店范围；单店趋势需要门店维度的按日汇总。')
])
variant('看板 C','变化来源','适合经营分析  对照单量 每单实收和渠道变化','dashboard-c',[
 ('适用角色与问题','老板、区域负责人和店长。了解收入波动与哪些指标同时变化，并决定下一步分析方向。'),
 ('页面顺序','筛选和三项核心指标 → 订单与每单实收双趋势 → 渠道金额对比 → 渠道占比及百分点变化。'),
 ('图表与选择理由','上下对齐的曲线分别显示订单数和平均每单实收，共用日期轴，单位为单、元/单。渠道分组柱对比本期与上期金额；两条占比条补充结构变化，避免双轴混淆。'),
 ('点击后的行为','点选同一日期联动两张趋势读数。渠道入口展示金额增减、占比和门店分布。渠道互斥、完整且口径一致时，显示对总净实收变化的金额贡献。'),
 ('具体使用情景','示例堂食减少 16,800 元，外卖增加 6,000 元，自提减少 1,200 元，合计减少 12,000 元。用户优先分析堂食；外卖占比增加 6 个百分点需与金额增长分开理解。'),
 ('数据支持与补充','已有日趋势、周期对比和渠道汇总可支撑图表。需核对渠道是否完整、退款是否一致分摊及对应订单分母。图表描述变化事实，不直接判定经营原因。')
])
variant('看板 D','经营日历热力图','新增方向  在长期日历中发现经营节奏','dashboard-d',[
 ('适用角色与问题','老板、区域负责人和店长。快速发现繁忙日、低谷日以及周末与工作日的重复模式。'),
 ('页面顺序','门店、90 天或近一年 → 周期核心指标 → 每日经营热度 → 选中日期汇总 → 当日渠道构成 → 周期观察。'),
 ('图表与选择理由','每格代表一天，按周排列，可切换订单量和净实收，单位为单/天或元/天。绿色深浅映射绝对数值；零值与缺失分别呈现。全年横向滑动，色阶入口说明阈值。'),
 ('点击后的行为','轻点日期查看当日数值及渠道汇总；长按拖动连续查看。放大月历及前后日按钮便于精确选择。周期对比使用完整上期，当日对比明确为上周同日。'),
 ('具体使用情景','用户看到多个周末色块更深，再查看周末与工作日日均单量，判断是否存在稳定的忙闲节奏。颜色表示数值大小，不表示盈利或经营好坏。'),
 ('数据支持与补充','已有按日数据可形成热力图。需确认 90 天或全年历史覆盖、每日完整性、营业日归属，以及支持筛选的门店和渠道日汇总；平均每单实收需对应有效订单数。')
])
variant('订单 A','订单总览','推荐默认订单页  面向每天数百至上千单的整体判断','order-a',[
 ('适用角色与问题','老板和店长。了解单量规模、相较上期的变化，以及订单来自哪些渠道和门店。'),
 ('页面顺序','日期、门店与更新 → 总订单及日均 → 已支付订单与平均每单实收 → 每日订单趋势 → 渠道结构 → 门店单量变化。'),
 ('图表与选择理由','每日单量用曲线，本期与上期逐日对应，单位为单。渠道占比条与各渠道单量并列，显示结构与规模；门店摘要按减少单量提示重点。'),
 ('点击后的行为','趋势仅查看日期汇总。点渠道查看该渠道周期表现；点门店查看单店汇总。门店、日期筛选同时影响所有数字和图表，不进入逐笔订单。'),
 ('具体使用情景','示例 7 天共有 8,580 单，日均 1,226 单，较上期增加 232 单；江湾店却减少 133 单。用户判断整体增长并不代表每家店增长，继续查看江湾的汇总表现。'),
 ('数据支持与补充','总订单、每日趋势及门店渠道汇总可支撑基础页。已支付订单及其占比需要支付状态汇总；平均每单实收需要核对净实收与已支付订单是否同口径。')
])
variant('订单 B','门店对比','适合区域负责人  区分单量规模与单量变化','order-b',[
 ('适用角色与问题','区域负责人和多店老板。比较门店订单规模，识别整体增长或下降由哪些门店贡献。'),
 ('页面顺序','全局筛选 → 所选门店订单及上期变化 → 门店订单对比图 → 门店分布汇总 → 统计口径。'),
 ('图表与选择理由','纵向细柱支持订单量与增减单量切换。订单量从零起算，按堂食、外卖、自提堆叠；增减单量以零基线区分正负。10 店示例横向滑动，不压缩成难以点选的柱。'),
 ('点击后的行为','点柱或门店汇总行查看单店周期指标、趋势与渠道构成；切换增减模式比较变化贡献。保留上期单量及变化率，不提供订单搜索。'),
 ('具体使用情景','滨江店 1,028 单、增加 109 单；江湾店 952 单、减少 133 单。用户可分别识别规模较大的门店和需要跟进的下降门店，不把低单量直接认定为经营差。'),
 ('数据支持与补充','门店订单数及周期对比可支持主图。堆叠柱需要门店与渠道交叉汇总，趋势需要门店日汇总；需标明新开、停业和数据不全门店对可比范围的影响。')
])
variant('订单 C','渠道与退款','适合店长与区域负责人  汇总观察渠道表现及退款情况','order-c',[
 ('适用角色与问题','店长和区域负责人。了解渠道单量变化，并判断退款指标是否需要进一步关注。'),
 ('页面顺序','筛选 → 总订单与退款订单比例 → 各渠道本期上期对比 → 退款与取消概况 → 每日退款趋势 → 统计口径。'),
 ('图表与选择理由','渠道分组柱对比订单绝对数量，单位单；同时显示占比、增减单量与退款比例。退款日趋势用曲线，支持与上期对照，便于区分单次波动与连续变化。'),
 ('点击后的行为','点渠道查看其订单、净实收及退款汇总；点日期查看当日统计。退款与取消分别呈现，不进入具体退款单或订单。'),
 ('具体使用情景','示例整体单量增长，但外卖减少 48 单，外卖退款订单比例为 4.0%。负责人继续比较该渠道的门店和上期基数，不能据此直接断言服务或产品存在问题。'),
 ('数据支持与补充','已有渠道订单数可支持分组柱。退款订单数、退款金额、已支付订单数和未付款取消数均需确认可用；缺少这些字段时应显示说明，隐藏相应分析。')
])

page();h('方案比较与推荐',1)
p('下表中的信息量和判断效率是设计评估，并非用户测试结果。七个版本分别强调不同问题，适合按角色或分析目的切换，无需把全部内容集中到首页。')
table(['版本','信息量','最有效的判断','额外数据依赖'],[
 ['看板 A','中','最快建立整体经营判断','较低  核对每单实收口径'],
 ['看板 B','中','最快确定门店跟进顺序','中  可比门店与日汇总'],
 ['看板 C','中高','定位与变化相关的指标和渠道','中高  统一渠道和日指标口径'],
 ['看板 D','中','识别长期忙闲与周期模式','中高  连续历史与缺失标记'],
 ['订单 A','中','最快了解全局单量与结构','中  支付状态汇总'],
 ['订单 B','中','比较门店规模和增减贡献','中  门店与渠道交叉汇总'],
 ['订单 C','较高','观察渠道差异和退款变化','高  完整退款与取消数据'],
],[.85,.65,2.85,2.85])
h('默认看板选择经营总览')
p('看板 A 同时给出净实收、单量和平均每单实收，再用趋势、门店和渠道引导下层查看，最符合打开应用后的快速判断目标。看板 B 可供区域负责人优先使用；看板 C 承接波动分析；看板 D 作为经营日历入口，补充长期节奏。')
h('默认订单页选择订单总览')
p('订单 A 以总量、趋势和结构为主，适合高单量业务。订单 B 承接门店差异，订单 C 承接渠道与退款汇总。三版均保持全局分析定位，页面中不增加逐笔查询入口。')
h('上线前的验证重点')
p('先让老板、区域负责人和店长分别完成一次 10 秒整体判断与 30 秒关注对象选择，观察是否需要解释图例、日期范围或指标口径。再验证 10 家以上门店的横向浏览、单店筛选、缺失数据和上期为零的状态。')
p('热力图需要额外验证颜色阈值是否容易理解，以及小色块是否难以选择。保留放大月历与前后日按钮，避免仅依赖精细点按。')

page();h('统一指标与交互要求',1)
h('指标和比较口径')
p('净实收表示收入结果，不等于利润；订单数不等于客流。平均每单实收以同范围净实收除以对应有效订单数计算，多店必须先合计金额和订单数，再求比值。不得直接平均各店均值。')
p('当前订单与日历模拟采用已支付订单作为每单实收分母，总订单可能同时包含未付款取消订单，因此不能直接用页面总订单替代分母。正式落地应统一有效订单、取消、全额退款和部分退款的处理规则，并在口径说明中展示。')
p('默认展示最近一个完整营业日或完整周期，只有确认支持实时数据后才使用今日实时。每次比较都显示本期及上期的明确日期范围；趋势首尾两天的差值不替代较上期变化。')
p('上期为零或缺失时不生成增长百分比。数据缺失不画成零；跨期尽量保持可比门店范围，说明新开、停业与不完整门店的影响。渠道不完整覆盖总额时说明覆盖范围，不强行补至 100%。')
h('订单与退款口径')
p('订单模拟中，总订单为已支付订单与未付款取消订单之和；已支付订单包含后续发生退款的订单。退款比例为发生退款的订单数除以已支付订单数，同一订单多次退款需要去重。退款金额仅计成功退款，包含部分与全额退款。')
p('退款模拟归属原订单营业日，并累计至数据更新时间；需与实际业务确认是否采用同一规则。退款订单、取消订单与成功退款金额应分别展示，不合并为一个未经定义的异常率。')
h('图表和触控')
p('手机参考宽度约 390px，使用分层卡片与等宽数字。日期和单选、多选、全选门店筛选同步作用于指标、图表及汇总。普通控件触控区域至少约 44px，图例保留文字和数值，减少装饰与动效。')
p('曲线必须通过真实采样点，避免过度平滑产生不存在的峰谷。轻点选择日期，长按拖动查看连续值，普通滑动移动日期窗口。每张图明确指标、单位、时间范围和比较对象。')
h('状态处理')
table(['状态','应呈现的内容'],[
 ['加载','保留筛选与卡片结构，用骨架占位，避免短暂显示 0。'],
 ['无数据','说明所选范围暂无数据，提供调整日期或门店入口。'],
 ['部分缺失','显示覆盖门店或日期范围，缺失位置保留断点或纹理；不可比时停用差异。'],
 ['请求失败','说明加载失败并提供重试；如保留旧数据，应标记原更新时间。'],
],[1.1,6.1])
compress=OxmlElement('w:doNotAutoCompressPictures');doc.settings.element.append(compress)
OUT.parent.mkdir(parents=True,exist_ok=True);doc.save(OUT)
print(OUT)
