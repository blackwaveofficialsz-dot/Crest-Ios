from pathlib import Path

from PIL import Image
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


BASE = Path('/Users/blackwave/Documents/Crest/work/mobile-ui-plan')
SOURCE = BASE / 'Crest-经营分析手机端UI规划.docx'
OUTPUT = BASE / 'Crest-经营分析手机端UI规划-补充订单详情.docx'
SCREENS = BASE / 'screens'
PURPLE = RGBColor(105, 77, 202)
INK = RGBColor(29, 30, 37)
MUTED = RGBColor(99, 101, 115)


def add_text(paragraph, value, size=9.2, bold=False, color=INK):
    run = paragraph.add_run(value)
    run.font.name = 'PingFang SC'
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'PingFang SC')
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return run


def shade(cell, fill):
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    cell._tc.get_or_add_tcPr().append(shd)


def crop_detail_screens():
    for mode in ('chart', 'table'):
        for theme in ('dark', 'light'):
            source = Image.open(SCREENS / f'order-detail-{mode}-{theme}-full.jpg').convert('RGB')
            center = source.width // 2
            source.crop((center - 201, 20, center + 201, source.height - 18)).save(
                SCREENS / f'order-detail-{mode}-{theme}.png', optimize=True
            )


doc = Document(SOURCE)
anchor = next(p._element for p in doc.paragraphs if p.text == '03 / 经营分析模块')
anchor = anchor.getprevious()  # Existing page-break paragraph before 产品
assert anchor.tag == qn('w:p') and anchor.xpath('.//w:br')
original_body = list(doc.element.body)
assert doc.paragraphs[17].text == '订单'


def label(value):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    add_text(p, value, 8, True, PURPLE)


def heading(value, subtitle):
    p = doc.add_paragraph(style='Heading 1')
    p.paragraph_format.space_after = Pt(3)
    p.add_run(value)
    q = doc.add_paragraph()
    q.paragraph_format.space_after = Pt(9)
    add_text(q, subtitle, 9.5, color=MUTED)


def note(name, value):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    add_text(p, name + '  ', 9.2, True, PURPLE)
    add_text(p, value)


def screenshots(mode):
    table = doc.add_table(rows=2, cols=2)
    table.autofit = False
    for i, theme in enumerate(('dark', 'light')):
        table.columns[i].width = Cm(8.55)
        caption = table.cell(0, i)
        caption.text = ''
        caption.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        shade(caption, 'F0EEF8' if theme == 'dark' else 'F4F4F7')
        p = caption.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        add_text(p, '深色模式' if theme == 'dark' else '浅色模式', 8.3, True)

        cell = table.cell(1, i)
        cell.text = ''
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(3)
        p.add_run().add_picture(
            str(SCREENS / f'order-detail-{mode}-{theme}.png'),
            width=Cm(7.0 if mode == 'table' else 7.4),
        )


crop_detail_screens()

doc.add_page_break()
label('02 / 订单 · 门店详情')
heading('点击门店后 · 折线图', '从门店排名进入独立详情，默认查看该店订单趋势。')
screenshots('chart')
note('进入方式', '点击任意门店排名行打开详情；顶部提供“返回门店排名”。')
note('趋势交互', '折线图与数据表切换；支持 7/30 日，并可沿线查看单日订单数及同期变化。')

doc.add_page_break()
label('02 / 订单 · 门店详情')
heading('点击门店后 · 数据表', '同一门店、同一日期周期下，切换查看逐日明细。')
screenshots('table')
note('表格内容', '按日期展示订单数、净实收、客单价与退款金额；窄屏横向滑动查看完整列。')
note('状态保持', '切换折线图/数据表时，保留已选门店与 7/30 日周期。')

new_elements = [element for element in doc.element.body if element not in original_body]
for element in new_elements:
    anchor.addprevious(element)

doc.save(OUTPUT)
print(OUTPUT)
