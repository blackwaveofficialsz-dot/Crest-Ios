from pathlib import Path

from PIL import Image
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


BASE = Path('/Users/blackwave/Documents/Crest/work/mobile-ui-plan')
SHOTS = BASE / 'screens'
OUT = BASE / 'Crest-经营分析手机端UI规划.docx'
PURPLE = RGBColor(105, 77, 202)
INK = RGBColor(29, 30, 37)
MUTED = RGBColor(99, 101, 115)
LINE = RGBColor(226, 226, 234)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tc_pr.append(shd)


def border_bottom(paragraph, color='E2E2EA', size='8'):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:color'), color)
    bottom.set(qn('w:sz'), size)
    borders.append(bottom)
    p_pr.append(borders)


def font(run, size=10, bold=False, color=INK):
    run.font.name = 'PingFang SC'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'PingFang SC')
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return run


def add_text(paragraph, text, size=10, bold=False, color=INK):
    return font(paragraph.add_run(text), size, bold, color)


def crop_screens():
    for module in ('dashboard', 'orders', 'products', 'labor', 'alerts'):
        for theme in ('dark', 'light'):
            source = Image.open(SHOTS / f'{module}-{theme}-full.png').convert('RGB')
            center = source.width // 2
            # Same 390 px phone in both captures; trim browser canvas while
            # retaining a narrow margin around its rounded silhouette.
            cropped = source.crop((center - 201, 21, center + 201, 931))
            cropped.save(SHOTS / f'{module}-{theme}.png', optimize=True)


def page_number(section):
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_text(footer, 'CREST  /  手机端经营分析 UI 规划    ·    ', 8, color=MUTED)
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)


doc = Document()
section = doc.sections[0]
section.page_width = Cm(21)
section.page_height = Cm(29.7)
section.top_margin = Cm(1.65)
section.bottom_margin = Cm(1.55)
section.left_margin = Cm(1.65)
section.right_margin = Cm(1.65)
section.header_distance = Cm(0.8)
section.footer_distance = Cm(0.8)
page_number(section)

normal = doc.styles['Normal']
normal.font.name = 'PingFang SC'
normal._element.rPr.rFonts.set(qn('w:eastAsia'), 'PingFang SC')
normal.font.size = Pt(9.5)
normal.font.color.rgb = INK
normal.paragraph_format.space_after = Pt(0)
normal.paragraph_format.line_spacing = 1.17

for name, size, color in [('Title', 28, INK), ('Heading 1', 17, INK), ('Heading 2', 11, INK)]:
    style = doc.styles[name]
    style.font.name = 'PingFang SC'
    style._element.rPr.rFonts.set(qn('w:eastAsia'), 'PingFang SC')
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = color


def label(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(5)
    add_text(p, text, 8, True, PURPLE)
    return p


def heading(title, subtitle=None):
    p = doc.add_paragraph(style='Heading 1')
    p.paragraph_format.space_after = Pt(3)
    p.add_run(title)
    if subtitle:
        q = doc.add_paragraph()
        q.paragraph_format.space_after = Pt(9)
        add_text(q, subtitle, 9.5, color=MUTED)


def compact_row(name, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(0)
    add_text(p, name + '  ', 9.2, True, PURPLE)
    add_text(p, text, 9.2)
    return p


def screenshot_pair(module):
    table = doc.add_table(rows=2, cols=2)
    table.autofit = False
    for i, theme in enumerate(('dark', 'light')):
        table.columns[i].width = Cm(8.55)
        cap = table.cell(0, i)
        cap.text = ''
        cap.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        shade(cap, 'F0EEF8' if theme == 'dark' else 'F4F4F7')
        p = cap.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        add_text(p, '深色模式' if theme == 'dark' else '浅色模式', 8.3, True, INK)

        cell = table.cell(1, i)
        cell.text = ''
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(0)
        p.add_run().add_picture(str(SHOTS / f'{module}-{theme}.png'), width=Cm(7.75))
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


crop_screens()

# Cover
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(38)
p.paragraph_format.space_after = Pt(13)
add_text(p, 'CREST  /  PRODUCT DESIGN', 9, True, PURPLE)

p = doc.add_paragraph(style='Title')
p.paragraph_format.space_after = Pt(7)
p.add_run('经营分析 · 手机端 UI 规划')
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(27)
add_text(p, '基于当前模拟页的界面结构、交互说明与深浅模式对照', 12, color=MUTED)
border_bottom(p, '694DCA', '20')

table = doc.add_table(rows=1, cols=3)
table.autofit = False
for i, (value, caption) in enumerate((('05', '业务栏目'), ('10', '主题对照截图'), ('01', '手机端模拟页'))):
    cell = table.cell(0, i)
    cell.width = Cm(5.7)
    cell.text = ''
    shade(cell, 'F4F2FA')
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(11)
    p.paragraph_format.space_after = Pt(2)
    add_text(p, value, 20, True, PURPLE)
    p = cell.add_paragraph()
    p.paragraph_format.space_after = Pt(11)
    add_text(p, caption, 8.8, color=MUTED)

doc.add_paragraph().paragraph_format.space_after = Pt(6)
heading('规划目标')
for text in (
    '用同一套移动端结构呈现经营分析的五个栏目，让使用者先识别异常，再定位门店与时间趋势。',
    '保留已确认的交互方向：总览区域紧凑、普通页面滚动；不做吸附效果。',
    '以当前模拟页面为视觉基线，后续接入真实数据时再确认统计口径与权限。',
):
    p = doc.add_paragraph(style='Normal')
    p.style = 'List Bullet'
    p.paragraph_format.space_after = Pt(4)
    add_text(p, text, 10)

doc.add_paragraph().paragraph_format.space_after = Pt(3)
heading('信息架构')
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(7)
add_text(p, '看板  →  订单  →  产品  →  人效  →  报警', 11, True, PURPLE)
compact_row('统一入口', '顶部栏目导航与统计日期。每个栏目独立承载自己的核心指标和分析区。')
compact_row('主题规范', '深色、浅色保持相同内容层级与交互；通过文字、图例与色彩共同表达数据。')
compact_row('文档边界', '截图与数字均取自模拟页面，仅说明方案，不代表正式产品数据或最终指标定义。')

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(25)
add_text(p, '版本  2026.10.08    ·    页面基线  Crest 经营分析手机模拟页', 8.5, color=MUTED)

modules = [
    {
        'key': 'dashboard', 'number': '01', 'title': '看板',
        'subtitle': '从当日营收进入同日门店对比、周期趋势与渠道构成。',
        'structure': '精简营收概览、营业渠道筛选、10 店纵向堆叠柱状图、门店周期折线图、堂食/外卖/自提结构。',
        'interaction': '柱图左右滑动并点选查看门店明细；趋势门店可单选、多选、全选；支持 7/30 日切换和沿线查看单日数据。',
        'planning': '同日比较与时间趋势使用同一统计日期；渠道筛选联动图表。上线前确认净实收口径与跨店汇总规则。',
    },
    {
        'key': 'orders', 'number': '02', 'title': '订单',
        'subtitle': '先看 10 店订单表现，再进入单店的订单趋势和每日明细。',
        'structure': '紧凑指标区展示订单数、净实收、客单价、退款金额；下方按订单数排列门店。总览不重复放订单趋势。',
        'interaction': '点击排名行直接打开该店详情，不在列表原地展开；详情支持折线图/数据表切换与 7/30 日周期。',
        'planning': '详情表按日期列出订单数、净实收、客单价与退款；折线图沿用看板的周期对比和逐点查看体验。',
    },
    {
        'key': 'products', 'number': '03', 'title': '产品',
        'subtitle': '将热销商品与单品销售走势关联起来。',
        'structure': '核心指标、近 7 日产品销售趋势、热销产品前 6 名以及当前选中商品的明细。',
        'interaction': '点击商品排行切换趋势对象；点击折线点查看该日销售额，选中商品信息同步更新。',
        'planning': '商品销售额与财务净实收分开表达；真实接入时需确定 SKU、退款和套餐拆分口径。',
    },
    {
        'key': 'labor', 'number': '04', 'title': '人效',
        'subtitle': '用工时投入与每小时产出判断排班效率。',
        'structure': '实际/计划工时、参与人数、小时产出；下方分别展示小时产出趋势及实际/计划工时趋势。',
        'interaction': '趋势图区分实际和计划；逐日条目可展开查看净实收、小时产出与目标。',
        'planning': '建议固定“小时产出＝净实收÷实际工时”的说明，并明确跨店合并及排班缺失值处理方式。',
    },
    {
        'key': 'alerts', 'number': '05', 'title': '报警',
        'subtitle': '优先发现偏离基准的门店，并给出可解释的查看入口。',
        'structure': '报警数量、上/下偏离概览、最大偏离；偏离率图按绝对值排序，下面是需关注门店清单。',
        'interaction': '正负偏离分列识别；点击清单条目展开实际营收、基准营收、偏离率与原因提示。',
        'planning': '真实报警阈值、基准周期与原因生成规则需要业务确认；突出可操作信息，避免仅用颜色区分风险。',
    },
]

for item in modules:
    doc.add_page_break()
    label(item['number'] + ' / 经营分析模块')
    heading(item['title'], item['subtitle'])
    screenshot_pair(item['key'])
    compact_row('页面构成', item['structure'])
    compact_row('当前交互', item['interaction'])
    compact_row('落地要点', item['planning'])

doc.add_page_break()
label('IMPLEMENTATION / 落地顺序')
heading('后续实现与验收', '建议按“统一容器 → 图表交互 → 真实数据 → 状态校验”的顺序推进。')
for title, body in (
    ('01  统一页面框架', '沿用五栏目导航、统计日期、紧凑总览和普通滚动；深浅模式逐屏核对信息层级与对比度。'),
    ('02  图表与下钻', '先实现看板门店柱图和趋势折线图，再复用交互模式到订单单店详情；补齐横向滑动、触点、空状态。'),
    ('03  数据口径与权限', '确认营收、订单、商品、人效、报警的时间粒度、门店范围和各指标计算方式；按账号权限限定可见门店。'),
    ('04  手机端验收', '至少验证常见手机宽度、深浅模式、日期切换、10 店横滑、单/多/全选、图表与数据表切换，以及异常/空数据。'),
):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(15)
    p.paragraph_format.space_after = Pt(6)
    add_text(p, title, 11, True, PURPLE)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    add_text(p, body, 10)

doc.add_paragraph().paragraph_format.space_after = Pt(11)
heading('本稿说明')
p = doc.add_paragraph()
add_text(p, '本稿用于确认手机端经营分析的界面方向。截图来自同一模拟页面的深色、浅色展示；所有数据均为演示数据，正式开发应以产品指标定义和真实接口为准。', 9.5, color=MUTED)

doc.save(OUT)
print(OUT)
