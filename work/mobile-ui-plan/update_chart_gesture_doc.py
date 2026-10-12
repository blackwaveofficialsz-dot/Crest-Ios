from pathlib import Path
from tempfile import NamedTemporaryFile
import os

from docx import Document


path = Path(__file__).with_name("Crest-经营分析手机端UI规划.docx")
document = Document(path)
replacements = {
    "柱图左右滑动并点选查看门店明细；趋势门店可单选、多选、全选。折线图默认展示最近 30 日：普通左右滑动切换可见日期范围，向右回看历史、向左返回；轻点选日期，长按折线后拖动才连续查看每日数据。也可切换 7 日。拖动不修改原始数据。":
    "柱图左右滑动并点选门店；趋势门店可单选、多选、全选。折线图默认 30 日，可切换 7 日；左右滑动回看历史，轻点选日期，长按折线后拖动逐日查看。拖动不改原始数据。",
    "折线图默认展示最近 30 日：普通左右滑动切换可见日期范围，向右回看历史、向左返回；轻点选日期，长按折线后拖动才连续查看每日订单数及同期变化。也可切换 7 日。拖动不修改原始数据。":
    "折线图默认 30 日，可切换 7 日；左右滑动回看历史，轻点选日期，长按折线后拖动逐日查看订单数及同期变化。拖动不改原始数据。",
}

changed = 0
for paragraph in document.paragraphs:
    for run in paragraph.runs:
        if run.text in replacements:
            run.text = replacements[run.text]
            changed += 1

if changed != len(replacements):
    raise RuntimeError(f"Expected {len(replacements)} local edits, found {changed}")

with NamedTemporaryFile(dir=path.parent, suffix=".docx", delete=False) as temporary:
    temporary_path = Path(temporary.name)
try:
    document.save(temporary_path)
    check = Document(temporary_path)
    if not all(any(new in paragraph.text for paragraph in check.paragraphs) for new in replacements.values()):
        raise RuntimeError("Saved document did not contain both updated descriptions")
    os.replace(temporary_path, path)
finally:
    temporary_path.unlink(missing_ok=True)
